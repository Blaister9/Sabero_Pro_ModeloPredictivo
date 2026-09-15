"""Copia de trabajo desde el paquete oficial A&A. No modifica la base aceptada."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from lxml import etree as E
from docx import Document
import hashlib, json, re

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'outputs/CONACIC2026'
BASE = ROOT / 'entrega_congreso/Saber_pro_paper_CONGRESO_10P_FINAL.docx'
TEMPLATE = OUT / 'plantillas_oficiales/Plantilla_A&A.docx'
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
W = '{'+NS['w']+'}'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
base_hash = sha(BASE)
assert base_hash == '5657e9af580c1c555c32128930b7c98cbeb3358cbdd511cd1b7f7ab706db7bcb'
assert (OUT/'AUDITORIA_CAMERA_READY.md').exists()

def xml(b): return E.fromstring(b)
def data(e): return E.tostring(e,xml_declaration=True,encoding='UTF-8',standalone=True)
def text(p): return ''.join(p.xpath('.//w:t/text()',namespaces=NS))
def replace_text(p,new):
    # Preserve paragraph and first run formatting; replace only this text slot.
    first = p.find(W+'r')
    props = deepcopy(first.find(W+'rPr')) if first is not None and first.find(W+'rPr') is not None else None
    for c in list(p):
        if c.tag != W+'pPr': p.remove(c)
    r=E.SubElement(p,W+'r')
    if props is not None:r.append(props)
    E.SubElement(r,W+'t').text=new

with ZipFile(TEMPLATE) as z: pkg={n:z.read(n) for n in z.namelist()}
with ZipFile(BASE) as z: src={n:z.read(n) for n in z.namelist()}
td=xml(pkg['word/document.xml']); tb=td.find(W+'body'); source=xml(src['word/document.xml']).find(W+'body')
tp=tb.findall(W+'p'); prototypes=[deepcopy(p) for p in tp]; section=deepcopy(tb.find(W+'sectPr'))
for c in list(tb):tb.remove(c)
nodes=[deepcopy(c) for c in source if c.tag!=W+'sectPr']
pars=[p for p in nodes if p.tag==W+'p']
# Accepted table was wider than the official text area. Fit its existing grid.
for table in [n for n in nodes if n.tag==W+'tbl']:
    widths=[2450,1350,1350,1300,2388]
    pr=table.find(W+'tblPr');pr.find(W+'tblW').set(W+'w',str(sum(widths)))
    pr.find(W+'tblLayout').set(W+'type','fixed')
    for col,width in zip(table.find(W+'tblGrid'),widths):col.set(W+'w',str(width))
    for row in table.findall(W+'tr'):
        for cell,width in zip(row.findall(W+'tc'),widths):
            tcw=cell.find(W+'tcPr').find(W+'tcW')
            if tcw is not None:tcw.set(W+'w',str(width));tcw.set(W+'type','dxa')
changes={
29:'Los archivos ICFES usan la columna MEDIDA_AGREGACION para discriminar el tipo de estadística de cada fila. Se identificaron tres capas utilizables: PUNTAJE_PRUEBA, PUNTAJE_GLOBAL y NIVEL_DESEMPEÑO_PRUEBA. Se unieron las pruebas con el puntaje global mediante inner join por año, institución y programa; los niveles se añadieron mediante left join incluyendo NOMBRE_PRUEBA. El resultado es 127.716 filas wide con el target disponible al 100 %. La Figura 1 ilustra la distribución del PROMEDIO_GLOBAL, las filas por año y las pruebas y NBCs más frecuentes.',
36:'Se emplearon 15 variables de entrada: rezagos lag_1 y lag_2 de las observaciones anteriores disponibles del PROMEDIO_GLOBAL y de cada prueba, tendencias, volatilidad histórica, logaritmo del número de evaluados, año y variables categóricas. Los tres modelos tabulares usan TargetEncoder; en los artefactos evaluados, su modo automático produjo codificación multiclase (135 valores del objetivo), con 417 columnas transformadas en total. El split externo usa 2020–2023 para entrenamiento (98.954 filas) y 2024 para prueba (28.762 filas).',
38:'Finalmente se evaluaron cuatro familias: (1) Ridge con α automático (RidgeCV); (2) Lasso con α automático (LassoCV); (3) LightGBM con 50 trials de búsqueda Optuna [4] y 188 árboles en el artefacto final (mejor iteración 168 + 20); (4) Transformer encoder con early stopping (época 53/200). Las métricas de evaluación son RMSE, MAE y R². La interpretabilidad del modelo final se obtiene con SHAP [6].',
100:'[20] P. J. L. Adeodato y R. L. C. Silva Filho, “Where to aim? Factors that influence the performance of Brazilian secondary schools,” en Proc. 13th International Conference on Educational Data Mining, 2020, pp. 545–549.'}
changes[56]=text(pars[56]).replace('Los programas de Salud concentran los errores más bajos','Algunas áreas de Salud presentan errores bajos')
changes[69]=text(pars[69]).replace('recoleción','recolección')
log=[]
for i,new in changes.items():
    log.append({'paragraph_base':i,'before':text(pars[i]),'after':new,'reason':'Corrección objetiva documentada en AUDITORIA_CAMERA_READY.md'})
    replace_text(pars[i],new)
# Replace authors and affiliation using actual template paragraph roles.
for idx,proto,content in [(3,3,'Edwin Santiago Paz Bedoya¹, Nathalia Orozco Morales¹'),(4,4,'¹UNIMINUTO Virtual – Bogotá, Colombia')]:
    p=deepcopy(prototypes[proto]);replace_text(p,content)
    for vertical in p.findall('.//'+W+'vertAlign'):
        vertical.getparent().remove(vertical)
    pp=p.find(W+'pPr')
    if pp is None:pp=E.SubElement(p,W+'pPr')
    sp=pp.find(W+'spacing')
    if sp is None:sp=E.SubElement(pp,W+'spacing')
    sp.set(W+'before','0');sp.set(W+'after','20')
    nodes[nodes.index(pars[idx])]=p
email=deepcopy(prototypes[6]);replace_text(email,'edwin.paz@uniminuto.edu, nathalia.orozco@uniminuto.edu')
nodes.insert(nodes.index(pars[5]),email)
# Reduce inherited empty spacer, not body typography or scientific content.
for p in [pars[1],pars[5]]:
    pp=p.find(W+'pPr')
    if pp is None:pp=E.SubElement(p,W+'pPr')
    spacing=pp.find(W+'spacing')
    if spacing is None:spacing=E.SubElement(pp,W+'spacing')
    spacing.set(W+'before','0');spacing.set(W+'after','0');spacing.set(W+'line','80');spacing.set(W+'lineRule','exact')
# Remap figure relationships into the template package, preserving the official logo and headers.
rels=xml(pkg['word/_rels/document.xml.rels']); sr=xml(src['word/_rels/document.xml.rels']); mapping={}
for r in sr:
    if r.get('Type','').endswith('/image'):
        old=r.get('Id'); new='rIdConacic'+str(len(mapping)+1); target=r.get('Target');dest='media/conacic_'+Path(target).name
        pkg['word/'+dest]=src['word/'+target];n=deepcopy(r);n.set('Id',new);n.set('Target',dest);rels.append(n);mapping[old]=new
for n in nodes:
    for e in n.iter():
        for a in list(e.attrib):
            if a.startswith('{'+NS['r']+'}') and e.get(a) in mapping:e.set(a,mapping[e.get(a)])
    tb.append(n)
pg=section.find(W+'pgNumType')
if pg is not None:pg.set(W+'start','1')
tb.append(section);pkg['word/document.xml']=data(td);pkg['word/_rels/document.xml.rels']=data(rels)
# Preserve official header geometry/logo; replace example editorial text and dates.
for name in list(pkg):
    if re.match(r'word/(header|footer)\d+\.xml$',name):
        e=xml(pkg[name]);ps=e.findall('.//'+W+'p')
        for p in ps:
            t=text(p)
            if 'Abstraction' in t:replace_text(p,'Abstraction & Application · CONACIC 2026 · DRAFT')
            elif 'Diseño' in t:replace_text(p,'Saber Pro · Paz Bedoya y Orozco Morales · DRAFT')
            elif 'Fecha de recepción' in t:replace_text(p,'DRAFT · Pendiente de análisis de similitud e IA y revisión final')
        if name=='word/header1.xml':
            p=ps[0]; pp=p.find(W+'pPr')
            if pp is None:pp=E.SubElement(p,W+'pPr')
            tabs=pp.find(W+'tabs')
            if tabs is not None:pp.remove(tabs)
            tabs=E.SubElement(pp,W+'tabs');tab=E.SubElement(tabs,W+'tab');tab.set(W+'val','right');tab.set(W+'pos','8838')
            r=E.SubElement(p,W+'r');E.SubElement(r,W+'tab')
            fld=E.SubElement(p,W+'fldSimple');fld.set(W+'instr','PAGE');r=E.SubElement(fld,W+'r');E.SubElement(r,W+'t').text='2'
        pkg[name]=data(e)
# Normalize running text only; preserve first-page logo and section geometry.
for name in ['word/header1.xml','word/header2.xml']:
    h=E.Element(W+'hdr',nsmap={'w':NS['w']});p=E.SubElement(h,W+'p');pp=E.SubElement(p,W+'pPr')
    tabs=E.SubElement(pp,W+'tabs');t=E.SubElement(tabs,W+'tab');t.set(W+'val','right');t.set(W+'pos','8838')
    spacing=E.SubElement(pp,W+'spacing');spacing.set(W+'after','0');spacing.set(W+'before','0')
    r=E.SubElement(p,W+'r');rp=E.SubElement(r,W+'rPr');fonts=E.SubElement(rp,W+'rFonts');fonts.set(W+'ascii','Calibri');fonts.set(W+'hAnsi','Calibri');E.SubElement(rp,W+'sz').set(W+'val','18')
    E.SubElement(r,W+'t').text='Saber Pro · Paz Bedoya y Orozco Morales · DRAFT'
    r=E.SubElement(p,W+'r');E.SubElement(r,W+'tab')
    fld=E.SubElement(p,W+'fldSimple');fld.set(W+'instr','PAGE');r=E.SubElement(fld,W+'r');rp=E.SubElement(r,W+'rPr');E.SubElement(rp,W+'sz').set(W+'val','18');E.SubElement(r,W+'t').text='2'
    pkg[name]=data(h)
settings=xml(pkg['word/settings.xml']);upd=settings.find(W+'updateFields')
if upd is None:upd=E.SubElement(settings,W+'updateFields')
upd.set(W+'val','true');pkg['word/settings.xml']=data(settings)
# Remove sample core metadata, retain other opaque template parts.
core=xml(pkg['docProps/core.xml'])
for n in core:
    local=E.QName(n).localname
    if local=='title':n.text=text(pars[0])
    elif local=='creator':n.text='Edwin Santiago Paz Bedoya; Nathalia Orozco Morales'
    elif local=='description':n.text='DRAFT. No enviar. Pendiente de análisis de similitud e IA y revisión final.'
pkg['docProps/core.xml']=data(core)
ct=xml(pkg['[Content_Types].xml']);existing={x.get('Extension') for x in ct}
for r in xml(src['[Content_Types].xml']):
    if r.get('Extension') and r.get('Extension') not in existing:ct.append(deepcopy(r));existing.add(r.get('Extension'))
pkg['[Content_Types].xml']=data(ct)
dest=OUT/'CAMERA_READY_DRAFT.docx'
with ZipFile(dest,'w',ZIP_DEFLATED) as z:
    for n,b in pkg.items():z.writestr(n,b)
assert sha(BASE)==base_hash
doc=Document(dest);assert len(doc.tables)==1 and len(doc.inline_shapes)==5
(OUT/'CAMBIOS_CAMERA_READY.json').write_text(json.dumps(log,ensure_ascii=False,indent=2),encoding='utf-8')
source_doc=Document(BASE)
inventory=['# Inventario de la base aceptada','',f'Base: `{BASE.relative_to(ROOT)}`',f'SHA-256: `{base_hash}`','', 'Transcripción para auditoría. No constituye otro artículo.','']
for i,p in enumerate(source_doc.paragraphs):
    if p.text:inventory.extend([f'**P{i}** {p.text}',''])
for i,t in enumerate(source_doc.tables):
    inventory.append(f'## Tabla {i+1}')
    for r in t.rows:inventory.append(' | '.join(c.text for c in r.cells))
(OUT/'INVENTARIO_BASE.md').write_text('\n'.join(inventory),encoding='utf-8')
preserved=[n for n,b in pkg.items() if n in src and b==src[n]]
print(dest, 'figuras',len(doc.inline_shapes),'tablas',len(doc.tables))
