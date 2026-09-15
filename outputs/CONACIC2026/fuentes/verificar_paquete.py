"""Verifica los contratos de la entrega DRAFT sin entrenar ni editar documentos."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from PIL import Image
import fitz,hashlib,json,re

ROOT=Path(__file__).resolve().parents[3];OUT=ROOT/'outputs/CONACIC2026'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
base=ROOT/'entrega_congreso/Saber_pro_paper_CONGRESO_10P_FINAL.docx'
assert sha(base)=='5657e9af580c1c555c32128930b7c98cbeb3358cbdd511cd1b7f7ab706db7bcb'
template=OUT/'plantillas_oficiales/Plantilla_A&A.docx'
assert sha(template)=='024d8159d545ec94355e115ae1bb43449ed88a184d1f957917abf930ebb16bb8'
doc=Document(OUT/'CAMERA_READY_DRAFT.docx');src=Document(base)
assert doc.paragraphs[0].text==src.paragraphs[0].text
assert len(doc.tables)==1 and len(doc.inline_shapes)==5
assert [[c.text for c in r.cells] for r in doc.tables[0].rows]==[[c.text for c in r.cells] for r in src.tables[0].rows]
text='\n'.join(p.text for p in doc.paragraphs)
for value in ['Edwin Santiago Paz Bedoya','Nathalia Orozco Morales','UNIMINUTO Virtual','edwin.paz@uniminuto.edu','nathalia.orozco@uniminuto.edu']:assert value in text
assert len(re.findall(r'^\[\d+\]',text,re.M))==26
for a,b in zip(doc.sections,Document(template).sections):
    for attr in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin']:assert getattr(a,attr)==getattr(b,attr)
with ZipFile(template) as t,ZipFile(OUT/'CAMERA_READY_DRAFT.docx') as f,ZipFile(base) as s:
    preserved=[]
    for n in t.namelist():
        if n in ['word/styles.xml','word/numbering.xml'] or n.startswith('word/theme/') or n.startswith('word/media/'):
            assert t.read(n)==f.read(n);preserved.append(n)
    source_media={hashlib.sha256(s.read(n)).hexdigest() for n in s.namelist() if n.startswith('word/media/')}
    final_media={hashlib.sha256(f.read(n)).hexdigest() for n in f.namelist() if n.startswith('word/media/')}
    assert source_media<=final_media
pdfs={}
for name in ['CAMERA_READY_DRAFT.pdf','PRESENTACION_CONACIC2026.pdf']:
    pdf=fitz.open(OUT/name);assert len(pdf)==10
    assert all('DRAFT' in page.get_text() for page in pdf)
    pdfs[name]={'pages':len(pdf),'sha256':sha(OUT/name),'draft_all_pages':True}
with ZipFile(OUT/'PRESENTACION_CONACIC2026.pptx') as z:
    slides=[n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)]
    assert len(slides)==10
    notes=[n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)];assert len(notes)==10
    assert b'<a:tbl' in z.read('ppt/slides/slide7.xml')
    assert any('chart' in n and n.endswith('.xml') for n in z.namelist())
    assert any(n.endswith('.xlsx') for n in z.namelist())
    px=E.fromstring(z.read('ppt/presentation.xml'));size=px.xpath('//*[local-name()="sldSz"]')[0];assert (size.get('cx'),size.get('cy'))==('12192000','6858000')
config=json.loads((OUT/'CRONOMETRAJE.json').read_text(encoding='utf-8'));assert len(config['slides'])==10 and config['total_seconds']==510 and config['total_words']==1162
for i in range(1,11):assert Image.open(OUT/f'escenas/{i:02d}.jpg').size==(1920,1080)
report={'status':'DRAFT_VERIFIED','scientific_approval':False,'original_docx_unchanged':sha(base),'template_docx_unchanged':sha(template),'template_parts_preserved':preserved,'scientific_media_preserved':True,'table_metrics_preserved':True,'author_fields_confirmed':True,'pdfs':pdfs,'slides':10,'notes':10,'native_metric_table':True,'native_coverage_chart_with_workbook':True,'script_words':1162,'planned_seconds':510,'real_voice_duration_measured':False}
(OUT/'VERIFICACION_PAQUETE.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps(report,indent=2,ensure_ascii=False))
